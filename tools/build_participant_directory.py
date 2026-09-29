"""Build the Firebase timer participant directory from an XLSX export."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from pathlib import Path

import openpyxl


def normalized(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).casefold().replace("ß", "ss")
    return "".join(character for character in value if character.isascii() and character.isalnum())


def display_name(value: object) -> str:
    text = " ".join(str(value or "").split())
    if "," not in text:
        return text
    last_name, first_name = (part.strip() for part in text.split(",", 1))
    return " ".join(part for part in (first_name, last_name) if part)


def event_year(value: object) -> int | None:
    match = re.search(r"(?:19|20)\d{2}", str(value or ""))
    return int(match.group()) if match else None


def full_birth_year(value: object, reference_year: int) -> int | None:
    try:
        short_year = int(float(value))
    except (TypeError, ValueError):
        return None
    year = short_year if short_year >= 1000 else (reference_year // 100) * 100 + short_year
    if year > reference_year:
        year -= 100
    return year if 1900 <= year <= reference_year else None


def build(source: Path) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    worksheet = next((sheet for sheet in workbook.worksheets if sheet.sheet_state == "veryHidden"), None)
    if worksheet is None:
        raise RuntimeError("No source sheet found.")

    source_years = [
        year
        for row in worksheet.iter_rows(values_only=True)
        if (year := event_year(row[10])) is not None
    ]
    if not source_years:
        raise RuntimeError("No event year found.")
    first_included_year = max(source_years) - 1

    latest: dict[tuple[str, int, str], tuple[tuple[float, int], str, str]] = {}
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
        gender = str(row[0] or "").strip().lower()
        name = display_name(row[1])
        year_of_event = event_year(row[10])
        if gender not in {"m", "w"} or not name or year_of_event is None or year_of_event < first_included_year:
            continue
        birth_year = full_birth_year(row[11], year_of_event)
        if birth_year is None:
            continue
        organization = " ".join(str(row[12] or "").split())
        date_value = float(row[9]) if isinstance(row[9], (int, float)) else 0.0
        search_name = normalized(name)
        key = (search_name, birth_year, gender)
        timestamp = (date_value, row_number)
        if key not in latest or timestamp > latest[key][0]:
            latest[key] = (timestamp, name, organization)

    records: list[dict[str, object]] = []
    used_ids: set[str] = set()
    for (search_name, birth_year, gender), (_, name, organization) in latest.items():
        candidate_id = hashlib.sha256(f"{search_name}|{birth_year}|{gender}".encode()).hexdigest()[:16]
        if candidate_id in used_ids:
            raise RuntimeError("Directory id collision.")
        used_ids.add(candidate_id)
        records.append({
            "id": candidate_id,
            "name": name,
            "birthYear": birth_year,
            "gender": gender,
            "organization": organization,
            "searchName": search_name,
        })
    records.sort(key=lambda item: (normalized(str(item["name"])), int(item["birthYear"]), str(item["gender"])))
    return records


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    records = build(arguments.source)
    arguments.output.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {len(records)} unique records.")


if __name__ == "__main__":
    main()
